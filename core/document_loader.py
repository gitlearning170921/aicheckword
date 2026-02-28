"""文档加载器：支持 PDF、Word、Excel、TXT、Markdown 等格式，以及压缩包自动解压"""

import shutil
import zipfile
import tempfile
import tarfile
from pathlib import Path
from typing import List, Optional, Tuple

from langchain.schema import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    UnstructuredExcelLoader,
    TextLoader,
    UnstructuredMarkdownLoader,
)

from config import settings

LOADER_MAP = {
    ".pdf": PyPDFLoader,
    ".docx": Docx2txtLoader,
    ".doc": Docx2txtLoader,
    ".xlsx": UnstructuredExcelLoader,
    ".xls": UnstructuredExcelLoader,
    ".txt": TextLoader,
    ".md": UnstructuredMarkdownLoader,
}

# 支持的压缩格式
ARCHIVE_EXTENSIONS = (".zip", ".tar", ".tgz", ".gz")
SUPPORTED_DOC_EXTENSIONS = tuple(LOADER_MAP.keys())


def _load_xlsx_with_openpyxl(path: Path) -> List[Document]:
    """使用 openpyxl 读取 xlsx（当 UnstructuredExcelLoader 失败时备用），避免路径/格式兼容问题。"""
    import openpyxl
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    docs = []
    try:
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append("\t".join(str(c) if c is not None else "" for c in row))
            text = "\n".join(rows).strip()
            if text:
                docs.append(Document(
                    page_content=text,
                    metadata={"source_file": path.name, "file_type": ".xlsx", "sheet": sheet.title},
                ))
    finally:
        wb.close()
    return docs if docs else [Document(page_content="(空表)", metadata={"source_file": path.name, "file_type": ".xlsx"})]


def load_single_file(file_path) -> List[Document]:
    """加载单个文件，返回 Document 列表。失败时抛出带完整原因的异常。"""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix not in LOADER_MAP:
        raise ValueError(f"不支持的文件格式: {suffix}，支持的格式: {list(LOADER_MAP.keys())}")

    loader_cls = LOADER_MAP[suffix]
    last_error = None

    # Excel：先尝试 UnstructuredExcelLoader，失败则用 openpyxl 备用（兼容中文路径、复杂表等）
    if suffix in (".xlsx", ".xls"):
        if suffix == ".xlsx":
            try:
                loader = loader_cls(str(path))
                docs = loader.load()
            except Exception as e:
                last_error = e
                try:
                    docs = _load_xlsx_with_openpyxl(path)
                except Exception as e2:
                    err_msg = f"Excel 加载失败: {last_error!s}\n备用解析(openpyxl)也失败: {e2!s}"
                    raise RuntimeError(err_msg) from last_error
        else:
            # .xls 仅用原 loader，无 openpyxl 备用
            try:
                loader = loader_cls(str(path))
                docs = loader.load()
            except Exception as e:
                raise RuntimeError(f"xls 文件加载失败: {e!s}") from e
    else:
        try:
            loader = loader_cls(str(path))
            docs = loader.load()
        except Exception as e:
            # 仅对纯文本类用 TextLoader 兜底，其它直接抛
            if suffix == ".txt":
                try:
                    loader = TextLoader(str(path), encoding="utf-8")
                    docs = loader.load()
                except Exception as e2:
                    raise RuntimeError(f"文本加载失败: {e!s}; 备用编码失败: {e2!s}") from e
            else:
                raise RuntimeError(f"加载失败: {e!s}") from e

    for doc in docs:
        doc.metadata.setdefault("source_file", path.name)
        doc.metadata.setdefault("file_type", suffix)

    return docs


def is_archive(path) -> bool:
    """判断是否为支持的压缩文件"""
    path = Path(path)
    name = path.name.lower()
    if name.endswith(".tar.gz") or name.endswith(".tgz"):
        return True
    return path.suffix.lower() in (".zip", ".tar", ".gz")


def extract_archive(archive_path) -> Tuple[str, List[Path]]:
    """解压到临时目录，返回 (临时目录路径, 文档路径列表)。调用方需在完成后删除临时目录。"""
    archive_path = Path(archive_path)
    temp_dir = tempfile.mkdtemp(prefix="aicheckword_")
    temp_path = Path(temp_dir)
    doc_files = []

    try:
        if archive_path.suffix.lower() == ".zip":
            with zipfile.ZipFile(archive_path, "r") as zf:
                for name in zf.namelist():
                    if name.endswith("/") or "/__MACOSX" in name:
                        continue
                    base = Path(name).name
                    if not base or base.startswith("."):
                        continue
                    zf.extract(name, temp_dir)
                    extracted = temp_path / name
                    if extracted.is_file():
                        suf = extracted.suffix.lower()
                        if suf in SUPPORTED_DOC_EXTENSIONS:
                            doc_files.append(extracted)

        elif archive_path.name.lower().endswith((".tar.gz", ".tgz")) or archive_path.suffix.lower() in (".tar", ".gz"):
            with tarfile.open(archive_path, "r:*") as tf:
                for member in tf.getmembers():
                    if not member.isfile():
                        continue
                    name = member.name
                    if "/__MACOSX" in name or name.startswith("."):
                        continue
                    tf.extract(member, temp_dir)
                    extracted = temp_path / member.name
                    if extracted.is_file():
                        suf = extracted.suffix.lower()
                        if suf in SUPPORTED_DOC_EXTENSIONS:
                            doc_files.append(extracted)
        else:
            shutil.rmtree(temp_dir, ignore_errors=True)
            raise ValueError("不支持的压缩格式: " + str(archive_path.suffix))

        return temp_dir, doc_files

    except Exception:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise


def load_directory(dir_path) -> List[Document]:
    """加载目录下所有支持格式的文件"""
    path = Path(dir_path)
    all_docs = []

    for ext in LOADER_MAP:
        for file_path in path.rglob(f"*{ext}"):
            try:
                docs = load_single_file(file_path)
                all_docs.extend(docs)
            except Exception as e:
                print(f"加载文件失败 {file_path}: {e}")

    return all_docs


def split_documents(
    documents: List[Document],
    chunk_size: Optional[int] = None,
    chunk_overlap: Optional[int] = None,
) -> List[Document]:
    """将文档分块"""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size or settings.chunk_size,
        chunk_overlap=chunk_overlap or settings.chunk_overlap,
        separators=["\n\n", "\n", "。", "；", "，", " ", ""],
    )
    return splitter.split_documents(documents)


def load_and_split(file_path) -> List[Document]:
    """加载并分块单个文件"""
    docs = load_single_file(file_path)
    return split_documents(docs)


def load_and_split_directory(dir_path) -> List[Document]:
    """加载并分块目录下所有文件"""
    docs = load_directory(dir_path)
    return split_documents(docs)
